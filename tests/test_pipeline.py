from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from wechat_digest.analyzer import analyze_rows
from wechat_digest.collector import normalize_wx_messages
from wechat_digest.config import AppConfig, StorageConfig
from wechat_digest.db import Database
from wechat_digest.report import write_report


class FakeLLM:
    def analyze(self, transcript):
        return {
            "issues": [
                {
                    "title": "升级后无法启动",
                    "category": "崩溃",
                    "priority": "P0",
                    "summary": "2.4.1 升级后无法启动",
                    "affected_version": "2.4.1",
                    "reproduction_steps": ["升级", "重启"],
                    "error_codes": ["E1042"],
                    "source_message_ids": [2],
                    "confidence": 0.95,
                }
            ]
        }


def test_collect_deduplicate_analyze_and_report(tmp_path: Path):
    config = AppConfig(
        group_name="测试群",
        force_keywords=("#重要",),
        storage=StorageConfig(tmp_path / "db.sqlite", tmp_path / "reports", tmp_path / "agent.log"),
    )
    now = datetime(2026, 9, 27, 12, 0, tzinfo=config.tz)
    messages = [
        SimpleNamespace(type="time", time="2026-09-27 10:00:00"),
        SimpleNamespace(type="text", attr="friend", sender="小王", content="#重要 安装包签名错误", id="a"),
        SimpleNamespace(type="text", attr="friend", sender="小李", content="2.4.1 升级后无法启动 E1042", id="b"),
    ]
    records = normalize_wx_messages(messages, config, now)
    db = Database(tmp_path / "db.sqlite")
    db.init_schema()
    assert db.insert_messages(records) == 2
    assert db.insert_messages(records) == 0
    rows = db.messages_between(
        "测试群",
        datetime(2026, 9, 27, 0, 0, tzinfo=config.tz),
        datetime(2026, 9, 28, 0, 0, tzinfo=config.tz),
    )
    issues = analyze_rows(rows, "2026-09-27", FakeLLM(), 24000)
    assert len(issues) == 2
    assert issues[0].forced is True
    assert issues[1].priority == "P0"
    path = write_report(issues, date(2026, 9, 27), tmp_path / "reports", message_count=2)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["重要问题", "关键词强制收录", "运行统计"]
    assert workbook["重要问题"].max_row == 3
    db.mark_run(date(2026, 9, 27), "rules_only", now)
    assert db.should_run(date(2026, 9, 27), now, retry_minutes=60) is True
    db.close()

