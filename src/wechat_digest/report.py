from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analyzer import Issue


HEADERS = [
    "日期", "优先级", "问题标题", "问题分类", "问题描述", "影响版本", "复现步骤",
    "反馈人数", "首次反馈时间", "最后反馈时间", "反馈成员", "相关错误码",
    "是否强制收录", "收录原因", "原始消息ID", "模型置信度",
]


def _issue_row(issue: Issue) -> list[object]:
    return [
        issue.date, issue.priority, issue.title, issue.category, issue.description,
        issue.affected_version, "\n".join(issue.reproduction_steps), issue.reporter_count,
        issue.first_seen, issue.last_seen, "、".join(issue.reporters),
        "、".join(issue.error_codes), "是" if issue.forced else "否",
        issue.inclusion_reason, ",".join(str(x) for x in issue.evidence_ids),
        issue.confidence if issue.confidence is not None else "",
    ]


def _format_sheet(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [12, 10, 30, 14, 48, 15, 38, 10, 22, 22, 24, 20, 14, 20, 22, 12]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_report(
    issues: list[Issue], report_date: date, report_dir: Path, *, message_count: int
) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    path = report_dir / f"{report_date.isoformat()}-群聊问题日报.xlsx"
    workbook = Workbook()
    important = workbook.active
    important.title = "重要问题"
    important.append(HEADERS)
    for issue in issues:
        important.append(_issue_row(issue))
    _format_sheet(important)

    forced = workbook.create_sheet("关键词强制收录")
    forced.append(HEADERS)
    for issue in issues:
        if issue.forced:
            forced.append(_issue_row(issue))
    _format_sheet(forced)

    stats = workbook.create_sheet("运行统计")
    stats.append(["项目", "值"])
    stats.append(["报告日期", report_date.isoformat()])
    stats.append(["消息总数", message_count])
    stats.append(["重要问题数", len(issues)])
    stats.append(["关键词强制收录数", sum(1 for issue in issues if issue.forced)])
    stats.column_dimensions["A"].width = 24
    stats.column_dimensions["B"].width = 48
    for cell in stats[1]:
        cell.font = Font(bold=True)

    workbook.save(path)
    return path

